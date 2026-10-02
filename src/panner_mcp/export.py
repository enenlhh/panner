"""Human-facing exports of screening results: a colour-coded xlsx workbook and two RIS files.

The export is a derived, read-only view of confirmed products (screening results,
the citation pool, and the frozen protocol); it writes into a fresh directory and
never mutates its inputs, so it runs in one step without the preview/confirm ritual.
"""

from __future__ import annotations

import json
import re
import shutil
import tempfile
from datetime import datetime, timezone
from pathlib import Path

from .evaluate import evaluate_rules
from .screening import _resolve_file
from .workspace import sha256_file, workspace_state

EXPORTS = ("screening-report.xlsx", "screening-retained.ris", "screening-excluded.ris")
STATUS_FILL = {
    "excluded": "FFE8DADA",  # light red
    "retained": "FFE2EFDA",  # light green
    "failed": "FFEDEDED",    # light grey
    "pending": "FFFFF2CC",   # light yellow
}
HEADER_FILL = "FFD9E1F2"


def _load_runs(workspace_path: str, results_path: str | None) -> tuple[Path, list[dict], list[dict], dict, dict]:
    workspace, state = workspace_state(workspace_path)
    relative = results_path or state.get("active_screening")
    if not relative:
        raise ValueError("this workspace has no confirmed screening results; run step 04 first")
    requested = Path(relative)
    results_file = (requested if requested.is_absolute() else workspace / requested).resolve(strict=True)
    if (not results_file.is_relative_to(workspace)
            or not re.fullmatch(r"screening/v[0-9]+/screening-results\.jsonl", str(results_file.relative_to(workspace)))):
        raise ValueError("results_path must identify a PANNER screening version in this workspace")
    manifest_file = results_file.parent / "manifest.json"
    if not results_file.is_file() or not manifest_file.is_file():
        raise ValueError("screening results or manifest are missing")
    manifest = json.loads(manifest_file.read_text(encoding="utf-8"))
    if manifest.get("results_sha256") != sha256_file(results_file):
        raise ValueError("screening results differ from their manifest")
    if manifest.get("judgments_sha256") != sha256_file(results_file.parent / "judgments.jsonl"):
        raise ValueError("screening judgments differ from their manifest")
    if manifest.get("run_config_snapshot_sha256") != sha256_file(results_file.parent / "run-config.json"):
        raise ValueError("screening run configuration differs from its snapshot")
    if not manifest.get("running", False) and manifest.get("events_sha256") != sha256_file(results_file.parent / "events.jsonl"):
        raise ValueError("screening event journal differs from its manifest")
    results = [json.loads(line) for line in results_file.read_text(encoding="utf-8").split("\n") if line.strip()]
    judgments = [json.loads(line) for line in (results_file.parent / "judgments.jsonl").read_text(encoding="utf-8").split("\n") if line.strip()]
    pool_file = Path(manifest["pool_path"])
    if manifest.get("pool_sha256") != sha256_file(pool_file):
        raise ValueError("citation pool differs from the screening manifest")
    if manifest.get("pool_manifest_sha256") != sha256_file(pool_file.parent / "dedup-manifest.json"):
        raise ValueError("citation pool manifest differs from the screening manifest")
    pool = [json.loads(line) for line in pool_file.read_text(encoding="utf-8").split("\n") if line.strip()]
    protocol_file = workspace / f"protocol/{manifest['protocol_version']}/protocol.json"
    if manifest.get("protocol_sha256") != sha256_file(protocol_file):
        raise ValueError("protocol differs from the screening manifest")
    if manifest.get("protocol_manifest_sha256") != sha256_file(protocol_file.parent / "manifest.json"):
        raise ValueError("protocol manifest differs from the screening manifest")
    protocol = json.loads(protocol_file.read_text(encoding="utf-8"))
    return workspace, results, judgments, manifest, {"pool": pool, "protocol": protocol, "results_file": results_file}


def _exclusion_basis(result: dict, protocol: dict, pool_by_id: dict[str, dict]) -> str:
    if result["status"] != "excluded":
        reasons = [item.get("reason", "") for item in result.get("undetermined", [])]
        return "; ".join(reason for reason in reasons if reason)
    clauses = {clause["id"]: clause for clause in protocol["criteria_map"]}
    parts = []
    for fired in result.get("basis", []):
        rule = next((r for r in protocol["decision_rules"] if r["id"] == fired["rule"]), None)
        quotes = []
        for ref in (rule or {}).get("criterion_refs", []):
            quote = ref.get("quote", "")
            if not quote.startswith("<"):
                quotes.append(quote)
        parts.append(f"[{fired['rule']}] " + " / ".join(quotes) if quotes else f"[{fired['rule']}]")
    duplicates = pool_by_id.get(result["citation_id"], {}).get("duplicate_of")
    if duplicates:
        parts.append(f"duplicate of {duplicates}")
    return " | ".join(parts)


def _ris_block(citation: dict, note: str | None) -> str:
    def one_line(value) -> str:
        return str(value).replace("\r", " ").replace("\n", " ").strip()

    def tag(name: str, value) -> list[str]:
        if value is None:
            return []
        if isinstance(value, list):
            return [f"{name}  - {one_line(item)}" for item in value if one_line(item)]
        text = one_line(value)
        return [f"{name}  - {text}"] if text else []

    lines = []
    ty = {"article": "JOUR", "inproceedings": "CONF", "inbook": "CHAP", "techreport": "RPRT"}.get(citation.get("entry_type", ""), "JOUR")
    lines.append(f"TY  - {ty}")
    lines.extend(tag("ID", citation["id"]))
    lines.extend(tag("TI", citation["title"]))
    lines.extend(tag("AU", (citation.get("authors") or "").split("; ")))
    lines.extend(tag("AB", citation["abstract"]))
    lines.extend(tag("JO", citation["journal"]))
    lines.extend(tag("JA", citation["container_title"]))
    lines.extend(tag("PY", citation["year"]))
    lines.extend(tag("VL", citation["volume"]))
    lines.extend(tag("IS", citation["issue"]))
    pages = citation.get("pages", "")
    if pages:
        lines.extend(tag("SP", pages))
    lines.extend(tag("DO", citation["doi"]))
    lines.extend(tag("AN", citation["pmid"]))
    lines.extend(tag("PT", citation.get("publication_types")))
    if note:
        lines.extend(tag("N1", note))
    lines.append("ER  - ")
    return "\n".join(lines)


def _write_ris(path: Path, rows: list[dict], pool_by_id: dict[str, dict], protocol: dict, with_notes: bool) -> None:
    blocks = []
    for row in rows:
        citation = pool_by_id.get(row["citation_id"])
        if citation is None:
            continue
        note = None
        if with_notes and row["status"] == "excluded":
            note = f"PANNER screening: excluded ({_exclusion_basis(row, protocol, pool_by_id)})"
        elif with_notes:
            note = f"PANNER screening: {row['status']}; send for further assessment"
        blocks.append(_ris_block(citation, note))
    path.write_text("\n".join(blocks) + ("\n" if blocks else ""), encoding="utf-8")


def export_screening(workspace_path: str, output_dir: str, results_path: str | None = None, *, experimental: bool = False) -> dict:
    """Generate the xlsx report and the retained/excluded RIS files from confirmed results."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workspace, results, judgments, manifest, context = _load_runs(workspace_path, results_path)
    if not experimental and (manifest.get("backend") != "calibrated" or manifest.get("run_config", {}).get("scope") is not None):
        raise ValueError("public export requires a complete calibrated-model screening run")
    pool_by_id = {row["id"]: row for row in context["pool"]}
    protocol = context["protocol"]
    judgment_by_id = {item["citation_id"]: item for item in judgments}
    signals = {signal["id"]: signal for signal in protocol["signals"]}
    expected_ids = {row["id"] for row in context["pool"] if row["duplicate_of"] is None}
    if experimental and manifest.get("run_config", {}).get("scope") is not None:
        expected_ids &= set(manifest["run_config"]["scope"])
    result_ids = [row["citation_id"] for row in results]
    if len(result_ids) != len(expected_ids) or set(result_ids) != expected_ids:
        raise ValueError("screening results do not account for every unique citation in the pool")
    for result in results:
        if result["status"] not in ("excluded", "retained"):
            continue
        judgment = judgment_by_id.get(result["citation_id"])
        if judgment is None:
            raise ValueError(f"screening judgment is missing for {result['citation_id']}")
        recalculated = evaluate_rules(protocol["decision_rules"], judgment)
        expected_status = "excluded" if recalculated["status"] == "excluded_candidate" else "retained"
        if (result["status"] != expected_status or result.get("basis") != recalculated["basis"]
                or result.get("undetermined") != recalculated["undetermined"]):
            raise ValueError(f"screening result is not reproducible for {result['citation_id']}")

    out = Path(output_dir).expanduser()
    if not out.is_absolute():
        raise ValueError("output_dir must be an absolute path")
    out = out.resolve()
    if out.exists() and any(out.iterdir()):
        raise ValueError("output_dir must be new or empty")
    out.mkdir(parents=True, exist_ok=True)

    results_sorted = sorted(results, key=lambda row: (
        row["status"] != "retained", row.get("review_priority") != "topic_unclear", row["citation_id"]
    ))

    # ---- xlsx ----
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary_rows = [
        ("Protocol version", manifest.get("protocol_version")),
        ("Protocol sha256", manifest.get("protocol_sha256")),
        ("Citation pool", manifest.get("pool_path")),
        ("Backend", manifest.get("backend")),
        ("Model", manifest.get("model")),
        ("Judge channel", manifest.get("judge_channel", "calibrated-decision-model")),
        ("Calls", manifest.get("calls")),
        ("Cost (USD)", manifest.get("cost_usd")),
        ("Generated at", datetime.now(timezone.utc).isoformat(timespec="seconds")),
        ("Excluded", manifest["counts"].get("excluded")),
        ("Retained", manifest["counts"].get("retained")),
        ("Failed", manifest["counts"].get("failed")),
        ("Pending", manifest["counts"].get("pending", 0)),
        ("Topic unclear", sum(row.get("review_priority") == "topic_unclear" for row in results)),
    ]
    summary.append(["PANNER screening report", None])
    summary["A1"].font = Font(bold=True, size=14)
    for key, value in summary_rows:
        summary.append([key, value])
        if isinstance(value, str):
            summary.cell(row=summary.max_row, column=2).data_type = "s"
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 80

    sheet = workbook.create_sheet("Citations")
    headers = ["id", "status", "review_priority", "title", "authors", "year", "journal", "doi", "pmid",
               "exclusion_basis", "source_file", "source_position"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
    for row in results_sorted:
        citation = pool_by_id.get(row["citation_id"], {})
        sheet.append([
            row["citation_id"], row["status"], row.get("review_priority", ""),
            citation.get("title", ""), citation.get("authors", ""), citation.get("year"),
            citation.get("journal", ""), citation.get("doi", ""), citation.get("pmid", ""),
            _exclusion_basis(row, protocol, pool_by_id),
            citation.get("source", {}).get("file", ""), citation.get("source", {}).get("position"),
        ])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
        fill = STATUS_FILL.get(row["status"])
        if fill:
            sheet.cell(row=sheet.max_row, column=2).fill = PatternFill("solid", fgColor=fill)
    widths = [10, 10, 18, 60, 32, 7, 26, 26, 13, 70, 34, 8]
    from openpyxl.utils import get_column_letter
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    detail = workbook.create_sheet("Signals")
    detail.append(["citation_id", "signal", "question", "on", "answer", "confidence"])
    for cell in detail[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
    for item in judgments:
        for answer in item["answers"]:
            signal = signals.get(answer["signal"], {})
            confidence = answer.get("confidence")
            detail.append([
                item["citation_id"], answer["signal"], signal.get("question", ""),
                answer.get("on", "record"), answer["answer"],
                confidence if confidence is not None else "",
            ])
            for cell in detail[detail.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    for index, width in enumerate([10, 8, 60, 10, 9, 11], 1):
        detail.column_dimensions[chr(64 + index)].width = width
    detail.freeze_panes = "A2"

    xlsx_path = out / EXPORTS[0]
    workbook.save(xlsx_path)

    # ---- RIS ----
    retained = [row for row in results_sorted if row["status"] != "excluded"]
    excluded = [row for row in results_sorted if row["status"] == "excluded"]
    retained_path = out / EXPORTS[1]
    excluded_path = out / EXPORTS[2]
    _write_ris(retained_path, retained, pool_by_id, protocol, with_notes=True)
    _write_ris(excluded_path, excluded, pool_by_id, protocol, with_notes=True)

    return {
        "status": "exported",
        "workspace_path": str(workspace),
        "output_dir": str(out),
        "files": {name: str(out / name) for name in EXPORTS},
        "counts": {
            "retained_ris": len(retained),
            "excluded_ris": len(excluded),
            "xlsx_rows": len(results_sorted),
        },
    }
